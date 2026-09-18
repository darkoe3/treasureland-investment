from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from core.importers import parse_daily_sheet_workbook
from core.models import AuditLog, Person, TPMCode, TerminalNumber
from core.tests.test_daily_sheet_imports import workbook_upload
from core.tests.test_phase_4 import Phase4Mixin
from core.tests import test_terminal_register as terminal_tests
from core.tests.test_terminal_register import upload
from core.tests import test_terminal_partial_import as partial_tests


class SubAgentRegisterTests(Phase4Mixin, APITestCase):
    setUp = terminal_tests.TerminalRegisterTests.setUp
    data = terminal_tests.TerminalRegisterTests.data
    create = terminal_tests.TerminalRegisterTests.create
    approve = partial_tests.PartialTerminalImportTests.approve
    counts = partial_tests.PartialTerminalImportTests.counts

    def preview(self, rows, policy="STRICT", mode="ONBOARD_MISSING", aliases=False):
        wb = Workbook()
        wb.active.append(["S/NOS", "SUB-AGENT NUMBER" if aliases else "SUB AGT NOS", "NAME" if aliases else "SUB NAME"])
        for row in rows:
            wb.active.append(row)
        stream = BytesIO()
        wb.save(stream)
        response = self.client.post(self.imports + "preview/", {
            "agency": self.musa.pk, "register_type": "SUB_AGENT", "mode": mode,
            "policy": policy, "file": SimpleUploadedFile("sub.xlsx", stream.getvalue())}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def test_three_columns_aliases_counts_ignored_and_metadata(self):
        batch = self.preview([[1, "000new", "New Owner"], [2, "another", self.person.full_name],
                              [3, None, None], [None, None, None]], aliases=True)
        payload = batch["preview_payload"]
        self.assertEqual(payload["ignored_blank_rows"], 2)
        self.assertEqual(payload["creation_counts"], {"people": 1, "sub_agent_numbers": 2, "terminals": 0})
        self.assertEqual(self.approve(batch).status_code, 200)
        self.assertFalse(TerminalNumber.objects.exists())
        self.assertEqual(TPMCode.objects.get(code="another").person, self.person)
        log = AuditLog.objects.get(action="TERMINAL_IMPORT_CONFIRMED")
        self.assertEqual(log.new_values["register_type"], "SUB_AGENT")
        for secret in ("New Owner", "000new", "sub.xlsx"):
            self.assertNotIn(secret, str(log.new_values))

    def test_strict_partial_field_errors_duplicates_and_conflicts(self):
        rows = [[1, "safe", "Safe"], [2, "missing-name", None], [3, None, "Missing code"],
                [4, "dupe", "First"], [5, "DUPE", "Second"],
                [6, self.tpm_a.code, "Wrong name"],
                [7, self.other_tpm.code, self.other_tpm.person.full_name]]
        strict = self.preview(rows)
        before = self.counts()
        self.assertEqual(self.approve(strict).status_code, 400)
        self.assertEqual(self.counts(), before)
        batch = self.preview(rows, policy="PARTIAL")
        parsed = batch["preview_payload"]["rows"]
        self.assertEqual(parsed[1]["reasons"], ["SUB NAME is required."])
        self.assertEqual(parsed[2]["reasons"], ["SUB AGT NOS is required."])
        self.assertEqual([r["classification"] for r in parsed[3:]],
                         ["DUPLICATE", "DUPLICATE", "NAME_CONFLICT", "AGENCY_CONFLICT"])
        self.assertEqual(self.approve(batch).status_code, 200)
        self.assertEqual(self.counts(), (before[0] + 1, before[1] + 1, before[2]))

    def test_link_permissions_acknowledgements_and_fixed_type(self):
        missing = self.preview([[1, "missing", "Missing"]], mode="LINK_EXISTING")
        self.assertEqual(self.approve(missing).status_code, 400)
        batch = self.preview([[1, self.tpm_a.code, self.person.full_name]], mode="LINK_EXISTING")
        for changes in ({"register_type": "TERMINAL"}, {"reason": ""}, {"acknowledged_counts": {}}):
            self.assertEqual(self.approve(batch, **changes).status_code, 400)
        self.assertEqual(self.approve(batch).status_code, 200)
        self.client.force_authenticate(self.acct)
        self.assertEqual(self.client.post(self.imports + "preview/", {"agency": self.musa.pk,
            "register_type": "SUB_AGENT", "mode": "ONBOARD_MISSING", "file": upload([])}, format="multipart").status_code, 403)

    def test_revalidation_and_atomic_rollback(self):
        batch = self.preview([[1, "one", "One"], [2, "two", "Two"]])
        before = self.counts()
        original = TPMCode.objects.create
        def fail_second(**kwargs):
            if kwargs["code"] == "two":
                raise ValidationError("Forced failure")
            return original(**kwargs)
        with patch("core.terminal_register.TPMCode.objects.create", side_effect=fail_second):
            self.assertEqual(self.approve(batch).status_code, 400)
        self.assertEqual(self.counts(), before)
        TPMCode.objects.create(code="ONE", person=self.other_tpm.person)
        self.assertEqual(self.approve(batch).status_code, 400)
        self.assertFalse(Person.objects.filter(full_name="Two").exists())

    def test_numeric_warning_invalid_name_and_template(self):
        batch = self.preview([[1, 123, "Valid"], [2, "bad", 123]], policy="PARTIAL")
        self.assertEqual(batch["warnings"][0]["row"], 2)
        self.assertTrue(batch["preview_payload"]["rows"][1]["excluded"])
        self.assertEqual(self.approve(batch, warnings_acknowledged=False).status_code, 400)
        response = self.client.get(self.imports + f"template/?agency={self.musa.pk}&register_type=SUB_AGENT")
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual([c.value for c in wb.active[1]], ["S/NOS", "SUB AGT NOS", "SUB NAME"])
        self.assertEqual(wb.active["B2"].number_format, "@")
        self.assertIn("No Terminal Numbers are created", str(list(wb["Instructions"].values)))

    def test_later_assignment_and_daily_sheet_without_terminal_preserves_history(self):
        sheet = self.sheet()
        txn = self.txn(sheet, self.tpm_a)
        before = (txn.person_name_snapshot, txn.tpm_code_snapshot, txn.terminal_number_snapshot)
        def parse():
            return parse_daily_sheet_workbook(workbook_upload(rows=[{"sub": self.tpm_a.code, "amounts": [10, 20]}],
                headers=["Premier", "Lucky"], register_rows=[(self.tpm_a.code, "", self.person.full_name)],
                raw_date=date(2026, 8, 24)), self.musa, date(2026, 8, 24))
        self.assertEqual(parse().errors, [])
        self.assertEqual(parse().payload["rows"][0]["terminal_number"], "")
        self.create()
        self.assertEqual(parse().errors, [])
        self.assertEqual(parse().payload["rows"][0]["terminal_number"], "")
        txn.refresh_from_db()
        self.assertEqual(before, (txn.person_name_snapshot, txn.tpm_code_snapshot, txn.terminal_number_snapshot))
        batch = terminal_tests.TerminalRegisterTests.preview(self, [[1, self.tpm_b.code, "later", self.person.full_name]])
        self.assertEqual(terminal_tests.TerminalRegisterTests.confirm(self, batch).status_code, 200)

    def test_daily_import_confirmation_does_not_infer_terminal(self):
        self.create()
        response = self.client.post("/api/daily-sheet-imports/preview/", {
            "agency": self.musa.pk, "transaction_date": "2026-08-24",
            "file": workbook_upload(rows=[{"sub": self.tpm_a.code, "amounts": [10, 20]}],
                headers=["Premier", "Lucky"], raw_date=date(2026, 8, 24))}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        result = self.client.post(f"/api/daily-sheet-imports/{response.data['id']}/confirm/", {}, format="json")
        self.assertEqual(result.status_code, 200, result.data)
        from core.models import TPMDailyTransaction
        self.assertEqual(TPMDailyTransaction.objects.get().terminal_number_snapshot, "")
