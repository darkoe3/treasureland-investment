"""Terminal master-data services. TPMCode remains the compatibility Sub-Agent model."""
import hashlib
from collections import Counter
from datetime import timedelta
from io import BytesIO

import openpyxl
from django.core.exceptions import ValidationError as ModelValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import filters, serializers, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response

from .importers import normalize_identifier, safe_filename, validate_workbook_bytes
from .models import Agency, AgentType, AuditLog, Person, TPMCode, TerminalImportBatch, TerminalNumber, User
from .permissions import IsSuperAdmin, SuperAdminOnlyWrites
from .serializers import AuditLogSerializer


def lock_register():
    # Stable parent locks also serialize empty slots (no terminal row exists yet).
    # Five agencies today; deterministic ordering prevents cross-agency deadlocks.
    list(Agency.objects.select_for_update().order_by("pk").values_list("pk", flat=True))


def identity(terminal):
    return {"terminal_number": terminal.terminal_number, "sub_agent_number": terminal.sub_agent_number_id,
            "sub_agent_number_value": terminal.sub_agent_number.code, "person": terminal.person_id,
            "person_name": terminal.person.full_name, "agency": terminal.agency_id,
            "agency_name": terminal.agency.name, "is_active": terminal.is_active}


def audit(user, obj, action_name, old=None, reason=""):
    is_batch = isinstance(obj, TerminalImportBatch)
    # Batch events contain metadata/counts only, never workbook rows.
    new = ({"status": obj.status, "policy": obj.preview_payload.get("policy", "STRICT"),
            "row_count": len(obj.preview_payload.get("rows", [])), "error_count": len(obj.errors),
            "warning_count": len(obj.warnings), "result_counts": obj.result_counts,
            "mode": obj.preview_payload.get("mode", "LINK_EXISTING"),
            "register_type": obj.preview_payload.get("register_type", "TERMINAL"),
            "creation_counts": obj.preview_payload.get("creation_counts", {}),
            "classification_counts": obj.preview_payload.get("summary", {}),
            "excluded_counts": obj.preview_payload.get("excluded_counts", {}),
            "imported_count": obj.result_counts.get("imported", obj.result_counts.get("created", 0) + obj.result_counts.get("unchanged", 0))} if is_batch else identity(obj))
    AuditLog.objects.create(user=user, agency=obj.agency, action=action_name,
                            model_name=type(obj).__name__, object_id=str(obj.pk),
                            old_values=old or {}, new_values=new, description=reason)


def resolve_owner(data):
    for field in ("agency", "person", "sub_agent_number"):
        if isinstance(data.get(field), bool) or not str(data.get(field, "")).isdigit():
            raise ValidationError({field: "Select a valid value."})
    try:
        code = TPMCode.objects.select_for_update().select_related("person__agency").get(pk=data["sub_agent_number"])
    except TPMCode.DoesNotExist:
        raise ValidationError({"sub_agent_number": "Select an existing Sub-Agent Number."})
    if code.person_id != int(data["person"]) or code.person.agency_id != int(data["agency"]):
        raise ValidationError("Person and agency must match the selected Sub-Agent Number.")
    return code


def terminal_text(value):
    if not isinstance(value, str) or not value.strip() or len(value.strip()) > 80:
        raise ValidationError({"terminal_number": "Enter a Terminal Number as text (1–80 characters)."})
    return value.strip()


def save_terminal(obj, user, action_name, old=None, reason=""):
    obj.updated_by = user
    obj.save()
    audit(user, obj, action_name, old, reason)
    return obj


def create_terminal(data, user, *, import_batch=False):
    code = resolve_owner(data)
    active = data.get("is_active", True)
    if not isinstance(active, bool):
        raise ValidationError({"is_active": "Use true or false."})
    obj = TerminalNumber(terminal_number=terminal_text(data.get("terminal_number")),
                         sub_agent_number=code, person=code.person, agency=code.person.agency,
                         is_active=active, created_by=user, updated_by=user,
                         deactivated_at=None if active else timezone.now(), deactivated_by=None if active else user)
    if import_batch:
        obj.save()
        return obj
    return save_terminal(obj, user, "TERMINAL_CREATED")


class TerminalSerializer(serializers.ModelSerializer):
    sub_agent_number_value = serializers.CharField(source="sub_agent_number.code", read_only=True)
    person_name = serializers.CharField(source="person.full_name", read_only=True)
    agency_name = serializers.CharField(source="agency.name", read_only=True)

    class Meta:
        model = TerminalNumber
        fields = "__all__"
        read_only_fields = [f.name for f in TerminalNumber._meta.fields]


class BatchSerializer(serializers.ModelSerializer):
    is_expired = serializers.SerializerMethodField()

    class Meta:
        model = TerminalImportBatch
        fields = "__all__"
        read_only_fields = [f.name for f in TerminalImportBatch._meta.fields]

    def get_is_expired(self, obj):
        return obj.expires_at <= timezone.now()


class SafeValidationMixin:
    def handle_exception(self, exc):
        if isinstance(exc, ModelValidationError):
            exc = ValidationError(getattr(exc, "message_dict", None) or exc.messages)
        elif isinstance(exc, IntegrityError):
            exc = ValidationError("An assignment changed or this Terminal Number already exists. Refresh and try again.")
        return super().handle_exception(exc)


class TerminalNumberViewSet(SafeValidationMixin, viewsets.GenericViewSet):
    permission_classes = [SuperAdminOnlyWrites]
    serializer_class = TerminalSerializer
    filter_backends = [filters.SearchFilter]
    search_fields = ["terminal_number", "sub_agent_number__code", "person__full_name", "agency__name"]

    def get_queryset(self):
        qs = TerminalNumber.objects.select_related("sub_agent_number", "person", "agency")
        if not self.request.user.is_super_admin:
            qs = qs.filter(agency__user_assignments__user=self.request.user)
        agency = self.request.query_params.get("agency")
        if agency:
            if not agency.isdigit():
                raise ValidationError({"agency": "Select a valid agency."})
            qs = qs.filter(agency_id=agency)
        active = self.request.query_params.get("active")
        if active in ("true", "false"):
            qs = qs.filter(is_active=active == "true")
        return qs

    def list(self, request):
        return Response(self.get_serializer(self.filter_queryset(self.get_queryset()), many=True).data)

    def retrieve(self, request, pk=None):
        return Response(self.get_serializer(self.get_object()).data)

    @transaction.atomic
    def create(self, request):
        lock_register()
        obj = create_terminal(request.data, request.user)
        return Response(self.get_serializer(obj).data, status=201)

    def locked_object(self):
        lock_register()
        return self.get_queryset().select_for_update().get(pk=self.get_object().pk)

    @transaction.atomic
    def partial_update(self, request, pk=None):
        obj = self.locked_object()
        if set(request.data) - {"terminal_number"}:
            raise ValidationError("Edit changes only the Terminal Number. Use the dedicated status or Reassign action for other changes.")
        old = identity(obj)
        obj.terminal_number = terminal_text(request.data.get("terminal_number"))
        save_terminal(obj, request.user, "TERMINAL_EDITED", old)
        return Response(self.get_serializer(obj).data)

    def change_status(self, request, active):
        obj = self.locked_object()
        if obj.is_active != active:
            old = identity(obj)
            obj.is_active = active
            obj.deactivated_at = None if active else timezone.now()
            obj.deactivated_by = None if active else request.user
            save_terminal(obj, request.user, "TERMINAL_REACTIVATED" if active else "TERMINAL_DEACTIVATED", old)
        return Response(self.get_serializer(obj).data)

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def deactivate(self, request, pk=None):
        return self.change_status(request, False)

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def reactivate(self, request, pk=None):
        return self.change_status(request, True)

    @action(detail=True, methods=["post"], permission_classes=[IsSuperAdmin])
    @transaction.atomic
    def reassign(self, request, pk=None):
        obj = self.locked_object()
        reason = request.data.get("reason")
        if not isinstance(reason, str) or not reason.strip() or len(reason) > 2000:
            raise ValidationError({"reason": "Enter a reason (1–2000 characters)."})
        if request.data.get("confirmed") is not True:
            raise ValidationError({"confirmed": "Explicit confirmation is required."})
        code = resolve_owner(request.data)
        if code.pk == obj.sub_agent_number_id:
            raise ValidationError("Choose a different Sub-Agent Number.")
        # An occupied target is never silently replaced: deactivate its terminal first.
        if TerminalNumber.objects.filter(sub_agent_number=code, is_active=True).exclude(pk=obj.pk).exists():
            raise ValidationError("Target already has an active terminal. Explicitly deactivate it before reassignment.")
        old = identity(obj)
        obj.sub_agent_number, obj.person, obj.agency = code, code.person, code.person.agency
        save_terminal(obj, request.user, "TERMINAL_REASSIGNED", old, reason.strip())
        return Response(self.get_serializer(obj).data)

    @action(detail=True, methods=["get"], permission_classes=[IsSuperAdmin])
    def history(self, request, pk=None):
        obj = self.get_object()
        logs = AuditLog.objects.filter(model_name="TerminalNumber", object_id=str(obj.pk))
        return Response(AuditLogSerializer(logs, many=True).data)


def normalize_name(value):
    return " ".join(str(value or "").split()).casefold()


def classify(code, number):
    existing = TerminalNumber.objects.filter(terminal_number__iexact=number).first()
    occupied = TerminalNumber.objects.filter(sub_agent_number=code, is_active=True).first()
    if existing and existing.sub_agent_number_id != code.id:
        return "Reassignment required", "Use Reassign with a reason and confirmation, then upload a fresh preview."
    if occupied and (not existing or occupied.pk != existing.pk):
        return "Conflict", "Sub-Agent Number already has an active terminal. Explicitly deactivate it before replacement."
    if existing:
        if not existing.is_active:
            return "Update required", "Reactivate the existing terminal explicitly, then upload a fresh preview."
        return "Unchanged", ""
    return "New", ""


CATEGORIES = ("CREATE_PERSON_SUBAGENT_TERMINAL", "ADD_SUBAGENT_TO_EXISTING_PERSON",
              "ADD_TERMINAL_TO_EXISTING_SUBAGENT", "UNCHANGED", "REASSIGNMENT_REQUIRED",
              "NAME_CONFLICT", "AGENCY_CONFLICT", "DUPLICATE", "INVALID", "INCOMPLETE")
NEW_CATEGORIES = CATEGORIES[:3]
SUB_AGENT_CATEGORIES = ("CREATE_PERSON_SUBAGENT", "ADD_SUBAGENT_TO_EXISTING_PERSON", "UNCHANGED",
                        "NAME_CONFLICT", "AGENCY_CONFLICT", "DUPLICATE", "INCOMPLETE",
                        "REASSIGNMENT_REQUIRED", "INVALID")
ACCEPTED_CATEGORIES = (*NEW_CATEGORIES, "CREATE_PERSON_SUBAGENT", "UNCHANGED")
REGISTER_HEADERS = {"TERMINAL": ["S/NOS", "SUB AGT NOS", "TERMINAL NOS", "NAME"],
                    "SUB_AGENT": ["S/NOS", "SUB AGT NOS", "SUB NAME"]}


def register_type_from(data):
    register_type = data.get("register_type", "TERMINAL")
    if register_type not in REGISTER_HEADERS:
        raise ValidationError({"register_type": "Select Terminal Register or Sub-Agent Register."})
    return register_type


def resolve_import_row(sub, number, name, agency, mode, register_type="TERMINAL"):
    codes = list(TPMCode.objects.select_related("person__agency").filter(code__iexact=sub))
    if len(codes) > 1:
        return "DUPLICATE", "Multiple case-insensitive Sub-Agent Numbers exist. Resolve manually.", None, None
    code = codes[0] if codes else None
    person = code.person if code else None
    category, message = "INVALID", ""
    terminal = TerminalNumber.objects.filter(terminal_number__iexact=number).first() if register_type == "TERMINAL" else None
    if (code and code.person.agency_id != agency.pk) or (terminal and terminal.agency_id != agency.pk):
        category, message = "AGENCY_CONFLICT", "Identifier belongs to another agency. Resolve manually."
    elif code and normalize_name(name) != normalize_name(person.full_name):
        category, message = "NAME_CONFLICT", "NAME does not match the database person."
    elif code and (not code.is_active or not person.is_active):
        message = "Sub-Agent Number and person must be active."
    elif terminal and (not code or terminal.sub_agent_number_id != code.pk):
        category, message = "REASSIGNMENT_REQUIRED", "Use the dedicated Reassign workflow."
    elif code and register_type == "SUB_AGENT":
        category = "UNCHANGED"
    elif code:
        state, message = classify(code, number)
        category = {"New": "ADD_TERMINAL_TO_EXISTING_SUBAGENT", "Unchanged": "UNCHANGED",
                    "Conflict": "REASSIGNMENT_REQUIRED", "Reassignment required": "REASSIGNMENT_REQUIRED",
                    "Update required": "REASSIGNMENT_REQUIRED"}[state]
    elif mode == "LINK_EXISTING":
        message = "Sub-Agent Number does not exist in the selected agency."
    else:
        matches = [p for p in Person.objects.filter(agency=agency) if normalize_name(p.full_name) == normalize_name(name)]
        if len(matches) > 1:
            category, message = "NAME_CONFLICT", "Multiple people have this normalized NAME. Manual resolution required."
        elif matches:
            person = matches[0]
            if not person.is_active:
                message = "Matching person is inactive. Resolve manually."
            else:
                category = "ADD_SUBAGENT_TO_EXISTING_PERSON"
        else:
            category = "CREATE_PERSON_SUBAGENT" if register_type == "SUB_AGENT" else "CREATE_PERSON_SUBAGENT_TERMINAL"
    if register_type == "SUB_AGENT":
        message = message.replace("NAME", "SUB NAME")
    return category, message, code.pk if code else None, person.pk if person else None


def creation_counts(rows, register_type="TERMINAL"):
    return {"people": len({normalize_name(r["name"]) for r in rows if r["classification"] in (NEW_CATEGORIES[0], "CREATE_PERSON_SUBAGENT")}),
            "sub_agent_numbers": sum(r["classification"] in (*NEW_CATEGORIES[:2], "CREATE_PERSON_SUBAGENT") for r in rows),
            "terminals": 0 if register_type == "SUB_AGENT" else sum(r["classification"] in NEW_CATEGORIES for r in rows)}


def parse_terminal_workbook(upload, agency, mode="LINK_EXISTING", policy="STRICT", register_type="TERMINAL"):
    sub_agent_only = register_type == "SUB_AGENT"
    headers = REGISTER_HEADERS[register_type]
    categories = SUB_AGENT_CATEGORIES if sub_agent_only else CATEGORIES
    data = validate_workbook_bytes(upload)
    errors, warnings, rows = [], [], []
    ignored_blank_rows = 0
    source_rows, source_formulas = [], []
    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationError("Workbook is corrupt or unsupported.") from exc
    try:
        if len(workbook.sheetnames) > 10:
            raise ValidationError("Workbook has too many worksheets.")
        sheet = workbook.worksheets[0]
        if sheet.max_row > 5001 or sheet.max_column > 20:
            raise ValidationError("Workbook exceeds the 5,000-row or 20-column limit.")
        actual_headers = [str(sheet.cell(1, col).value or "").strip().upper() for col in range(1, len(headers) + 1)]
        if sub_agent_only:
            actual_headers = [{"SUB-AGENT NUMBER": "SUB AGT NOS", "NAME": "SUB NAME"}.get(h, h) for h in actual_headers]
            if sheet.max_column > 3 and any(cell.value not in (None, "")
                    for row in sheet.iter_rows(min_col=4) for cell in row):
                raise ValidationError("Sub-Agent Register accepts only columns A-C: S/NOS, SUB AGT NOS, SUB NAME.")
        if actual_headers != headers:
            raise ValidationError(f"Row 1 must contain {', '.join(headers)} in columns A-{'C' if sub_agent_only else 'D'}.")
        seen_sub, seen_terminal = set(), set()
        for n, cells in enumerate(sheet.iter_rows(min_row=2, min_col=2, max_col=len(headers)), 2):
            source_rows.append([cell.value if isinstance(cell.value, (str, int, float, bool, type(None))) else str(cell.value) for cell in cells])
            source_formulas.append([cell.data_type == "f" for cell in cells])
            # Only B–D determine whether this is a data row. S/NOS is never required.
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                ignored_blank_rows += 1
                continue
            row_errors, row_warnings = [], []
            if any(cell.data_type == "f" for cell in cells):
                row_errors.append("Formula cells are not allowed in columns B–D.")
                values = ["" if c.data_type == "f" else c.value for c in cells]
            else:
                values = [c.value for c in cells]
            if any(isinstance(value, bool) for value in values):
                row_errors.append("Boolean cells are not valid identifiers or names.")
            sub = normalize_identifier(values[0], f"B{n}", row_warnings)
            number = "" if sub_agent_only else normalize_identifier(values[1], f"C{n}", row_warnings)
            raw_name = values[1] if sub_agent_only else values[2]
            name = str(raw_name or "").strip()
            if sub_agent_only and raw_name not in (None, "") and not isinstance(raw_name, str):
                row_errors.append("SUB NAME must be valid text.")
            supplied_name = name
            if not sub_agent_only and not name and sub and policy == "PARTIAL":
                owners = list(TPMCode.objects.select_related("person").filter(
                    code__iexact=sub, person__agency=agency))
                if len(owners) == 1 and owners[0].person_id:
                    name = owners[0].person.full_name
            required_fields = (("SUB AGT NOS", sub), ("SUB NAME", name)) if sub_agent_only else (
                ("SUB AGT NOS", sub), ("TERMINAL NOS", number), ("NAME", name))
            missing_fields = [label for label, value in required_fields if not value]
            row_errors.extend(f"{field} is required." for field in missing_fields)
            if len(sub) > 80 or len(number) > 80 or len(name) > 255:
                row_errors.append("Identifier or name exceeds the allowed length.")
            for value, seen, label in ((sub, seen_sub, "Sub-Agent Number"), (number, seen_terminal, "Terminal Number")):
                if value.casefold() in seen:
                    row_errors.append(f"Duplicate {label} in workbook.")
                if value:
                    seen.add(value.casefold())
            classification, code_id, person_id = "INCOMPLETE" if missing_fields and (policy == "PARTIAL" or sub_agent_only) else "INVALID", None, None
            if not row_errors:
                classification, message, code_id, person_id = resolve_import_row(sub, number, name, agency, mode, register_type)
                if message:
                    row_errors.append(message)
            elif any(message.startswith("Duplicate") for message in row_errors):
                classification = "DUPLICATE"
            rows.append({"row": n, "sub_agent_number": code_id, "sub_agent_number_value": sub,
                         "terminal_number": number, "person": person_id,
                         "name": supplied_name, "resolved_name": name, "classification": classification,
                         "missing_fields": missing_fields, "reasons": row_errors, "supplied_values": source_rows[-1],
                         "excluded": classification not in ACCEPTED_CATEGORIES})
            errors.extend({"row": n, "category": classification, "detail": message,
                           "message": f"Row {n}: {message}"} for message in row_errors)
            for warning in row_warnings:
                field = "SUB AGT NOS" if warning["cell"].startswith("B") else "TERMINAL NOS"
                warnings.append({"row": n, "field": field,
                                 "message": f"Row {n} — {field} is numeric and may have lost leading zeroes."})
        if policy == "PARTIAL" or sub_agent_only:
            # Every occurrence is unsafe: never choose an owner by workbook order.
            duplicate_fields = (("sub_agent_number_value", "Sub-Agent Number"),
                                ("terminal_number", "Terminal Number"))
            frequencies = {field: Counter(r[field].casefold() for r in rows if r[field])
                           for field, _ in duplicate_fields}
            for row in rows:
                duplicates = [f"Duplicate {label} in workbook." for field, label in duplicate_fields
                              if row[field] and frequencies[field][row[field].casefold()] > 1]
                if duplicates:
                    row["classification"], row["excluded"] = "DUPLICATE", True
                    row["reasons"].extend(message for message in duplicates if message not in row["reasons"])
            errors = [{"row": row["row"], "category": row["classification"], "detail": message,
                       "message": f"Row {row['row']}: {message}"}
                      for row in rows for message in row["reasons"]]
        if not rows:
            errors.append({"row": 2, "message": "Row 2: Workbook contains no completed rows."})
    finally:
        workbook.close()
    return data, {"mode": mode, "policy": policy, "source_rows": source_rows, "source_formulas": source_formulas,
                  "valid_count": sum(not r["excluded"] for r in rows),
                  "excluded_count": sum(r["excluded"] for r in rows),
                  "excluded_counts": {c: sum(r["excluded"] and r["classification"] == c for r in rows) for c in categories if c not in ACCEPTED_CATEGORIES},
                  "rows": rows, "ignored_blank_rows": ignored_blank_rows,
                  "creation_counts": creation_counts(rows, register_type),
                  "summary": {category: sum(r["classification"] == category for r in rows) for category in categories}, "register_type": register_type}, warnings, errors


class TerminalImportViewSet(SafeValidationMixin, viewsets.GenericViewSet):
    permission_classes = [IsSuperAdmin]
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    serializer_class = BatchSerializer

    def get_queryset(self):
        # Even Super Admins can only confirm/cancel their own previews.
        return TerminalImportBatch.objects.filter(uploader=self.request.user)

    def agency(self, data):
        if not str(data.get("agency", "")).isdigit():
            raise ValidationError({"agency": "Select an agency."})
        try:
            return Agency.objects.get(pk=data["agency"], is_active=True)
        except Agency.DoesNotExist:
            raise ValidationError({"agency": "Select an active agency."})

    def retrieve(self, request, pk=None):
        return Response(self.get_serializer(self.get_object()).data)

    @action(detail=False, methods=["get"])
    def template(self, request):
        self.agency(request.query_params)
        wb = openpyxl.Workbook()
        sheet = wb.active
        register_type = register_type_from(request.query_params)
        sub_agent_only = register_type == "SUB_AGENT"
        sheet.title = "Sub-Agent Register" if sub_agent_only else "Terminal Register"
        sheet.append(REGISTER_HEADERS[register_type])
        for col in (("B",) if sub_agent_only else ("B", "C")):
            sheet.column_dimensions[col].number_format = "@"
        for n in range(2, 5002):
            for col in ((2,) if sub_agent_only else (2, 3)):
                cell = sheet.cell(n, col, "")
                cell.data_type, cell.number_format = "s", "@"
        for col, width in ((("A", 10), ("B", 25), ("C", 35)) if sub_agent_only else (("A", 10), ("B", 25), ("C", 25), ("D", 35))):
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A2"
        instructions = wb.create_sheet("Instructions")
        for text in ("Choose the same agency when uploading. Data starts on row 2 of Terminal Register.",
                     "S/NOS is ignored. LINK_EXISTING requires existing people and Sub-Agent Numbers.",
                     "Keep identifiers as Text to preserve leading zeroes. Do not paste numeric cells.",
                     "ONBOARD_MISSING can create missing people and Sub-Agent Numbers after explicit confirmation. Existing names must match.",
                     "All three fields B–D are required. No formulas, macros or external links.",
                     "Preview first. Resolve conflicts using Edit, Deactivate, Reactivate or Reassign, then upload again."):
            if sub_agent_only:
                text = text.replace("Terminal Register", "Sub-Agent Register").replace(
                    "All three fields B–D are required.", "SUB AGT NOS and SUB NAME are required. No Terminal Numbers are created.")
            instructions.append([text])
        instructions.column_dimensions["A"].width = 110
        buffer = BytesIO()
        wb.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{"sub-agent" if sub_agent_only else "terminal"}-register.xlsx"'
        return response

    @action(detail=False, methods=["post"])
    @transaction.atomic
    def preview(self, request):
        agency = self.agency(request.data)
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": "Upload an .xlsx workbook."})
        mode = request.data.get("mode", "LINK_EXISTING")
        if mode not in ("LINK_EXISTING", "ONBOARD_MISSING"):
            raise ValidationError({"mode": "Select a valid import mode."})
        policy = request.data.get("policy", "STRICT")
        if policy not in ("STRICT", "PARTIAL") or (policy == "PARTIAL" and mode != "ONBOARD_MISSING"):
            raise ValidationError({"policy": "Partial importing requires Super Admin ONBOARD_MISSING mode."})
        register_type = register_type_from(request.data)
        data, payload, warnings, errors = parse_terminal_workbook(upload, agency, mode, policy, register_type)
        batch = TerminalImportBatch.objects.create(uploader=request.user, agency=agency,
            original_filename=safe_filename(upload.name.replace("\\", "/")), file_hash=hashlib.sha256(data).hexdigest(),
            preview_payload=payload, warnings=warnings, errors=errors, expires_at=timezone.now() + timedelta(hours=1))
        audit(request.user, batch, "TERMINAL_IMPORT_PREVIEWED")
        return Response(self.get_serializer(batch).data, status=201)

    def locked_batch(self):
        batch = self.get_queryset().select_for_update().get(pk=self.get_object().pk)
        if batch.status != "PREVIEWED":
            raise ValidationError("Only a previewed batch may be confirmed or cancelled.")
        if batch.expires_at <= timezone.now():
            raise ValidationError("This preview expired. Upload a fresh workbook.")
        return batch

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def confirm(self, request, pk=None):
        lock_register()
        batch = self.locked_batch()
        actor = User.objects.select_for_update().get(pk=request.user.pk)
        if not actor.is_active or not actor.is_super_admin:
            raise ValidationError("Super Admin permission is required.")
        if request.data.get("confirmed") is not True:
            raise ValidationError("Explicit confirmation is required.")
        register_type = batch.preview_payload.get("register_type", "TERMINAL")
        if "register_type" in request.data and request.data["register_type"] != register_type:
            raise ValidationError("The previewed register type cannot be changed. Upload a fresh preview.")
        policy = batch.preview_payload.get("policy", "STRICT")
        if "policy" in request.data and request.data["policy"] != policy:
            raise ValidationError("The previewed policy cannot be changed. Upload a fresh preview.")
        if batch.errors and policy == "STRICT":
            raise ValidationError("Resolve all blocking errors and create a fresh preview.")
        if policy == "PARTIAL":
            if batch.preview_payload.get("mode") != "ONBOARD_MISSING" or not request.user.is_super_admin:
                raise ValidationError("Partial importing requires Super Admin onboarding.")
            if not batch.preview_payload["valid_count"]:
                raise ValidationError("At least one valid row is required.")
            if request.data.get("exclusions_acknowledged") is not True:
                raise ValidationError("Explicitly acknowledge excluded and valid rows.")
            for field in ("valid_count", "excluded_count"):
                expected = request.data.get("expected_" + field)
                if type(expected) is not int or expected != batch.preview_payload[field]:
                    raise ValidationError(f"Expected {field} does not match preview.")
        if not Agency.objects.select_for_update().get(pk=batch.agency_id).is_active:
            raise ValidationError("Agency is inactive.")
        mode = batch.preview_payload.get("mode", "LINK_EXISTING")
        reason = request.data.get("reason", "")
        if batch.warnings and request.data.get("warnings_acknowledged") is not True:
            raise ValidationError("Verify numeric identifiers against the original source and acknowledge warnings.")
        if mode == "ONBOARD_MISSING" or register_type == "SUB_AGENT":
            if not isinstance(reason, str) or not reason.strip() or len(reason) > 2000:
                raise ValidationError("Enter an onboarding reason (1–2000 characters).")
            if request.data.get("acknowledged_counts") != batch.preview_payload["creation_counts"]:
                raise ValidationError("Explicitly acknowledge the displayed creation counts.")
        # Lock existing owners as well as parent agencies; validate the complete plan before writes.
        list(Person.objects.select_for_update().filter(agency=batch.agency).order_by("pk"))
        list(TPMCode.objects.select_for_update().order_by("pk"))
        list(TerminalNumber.objects.select_for_update().order_by("pk"))
        if "source_rows" in batch.preview_payload:
            wb = openpyxl.Workbook()
            wb.active.append(REGISTER_HEADERS[register_type])
            for index, values in enumerate(batch.preview_payload["source_rows"]):
                wb.active.append([None, *values])
                for column, value in enumerate(values, 2):
                    if isinstance(value, str):
                        is_formula = batch.preview_payload["source_formulas"][index][column - 2]
                        wb.active.cell(index + 2, column).data_type = "f" if is_formula else "s"
            stream = BytesIO()
            wb.save(stream)
            _, current, _, _ = parse_terminal_workbook(
                SimpleUploadedFile("revalidate.xlsx", stream.getvalue()), batch.agency, mode, policy, register_type)
            # Older Terminal Register previews predate the explicit type field.
            if "register_type" not in batch.preview_payload:
                current.pop("register_type", None)
            if current != batch.preview_payload:
                raise ValidationError("Rows, counts or assignments changed after preview. Upload again.")
        else:
            for row in batch.preview_payload["rows"]:
                category, message, code_id, person_id = resolve_import_row(
                    row["sub_agent_number_value"], row["terminal_number"], row["name"], batch.agency, mode, register_type)
                if message or (category, code_id, person_id) != (row["classification"], row["sub_agent_number"], row["person"]):
                    raise ValidationError("Identity or assignment changed after preview. Upload again.")
        counts = {"created": 0, "unchanged": 0}
        new_people = {}
        for row in batch.preview_payload["rows"]:
            if row.get("excluded"):
                continue
            if row["classification"] == "UNCHANGED":
                counts["unchanged"] += 1
                continue
            person_id, code_id = row["person"], row["sub_agent_number"]
            if not person_id:
                if not normalize_name(row["name"]):
                    raise ValidationError("NAME is required to create a Person.")
                key = normalize_name(row["name"])
                if key not in new_people:
                    new_people[key] = Person.objects.create(agency=batch.agency,
                        full_name=" ".join(row["name"].split()), agent_type=AgentType.SUBAGENT).pk
                person_id = new_people[key]
            if not code_id:
                code_id = TPMCode.objects.create(person_id=person_id, code=row["sub_agent_number_value"]).pk
            if register_type == "TERMINAL":
                create_terminal({**row, "person": person_id, "sub_agent_number": code_id, "agency": batch.agency_id}, request.user, import_batch=True)
            counts["created"] += 1
        if mode == "ONBOARD_MISSING" or register_type == "SUB_AGENT":
            counts.update(batch.preview_payload["creation_counts"])
        if policy == "PARTIAL" or register_type == "SUB_AGENT":
            counts.update(imported=batch.preview_payload["valid_count"],
                          excluded=batch.preview_payload["excluded_count"],
                          excluded_counts=batch.preview_payload["excluded_counts"])
        batch.status, batch.confirmed_at, batch.result_counts = "CONFIRMED", timezone.now(), counts
        batch.save()
        audit(request.user, batch, "TERMINAL_IMPORT_CONFIRMED", reason=reason.strip() if mode == "ONBOARD_MISSING" or register_type == "SUB_AGENT" else "")
        return Response(self.get_serializer(batch).data)

    @action(detail=True, methods=["get"], url_path="exclusions")
    def exclusions(self, request, pk=None):
        batch = self.get_object()
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Excluded rows"
        register_type = batch.preview_payload.get("register_type", "TERMINAL")
        sheet.append(["Excel row", *REGISTER_HEADERS[register_type][1:], "Category", "Reason"])
        for row in batch.preview_payload["rows"]:
            if row.get("excluded", row["classification"] not in ACCEPTED_CATEGORIES):
                supplied = batch.preview_payload.get("source_rows", [])[row["row"] - 2] if "source_rows" in batch.preview_payload else [row["sub_agent_number_value"], row["terminal_number"], row["name"]]
                reasons = row.get("reasons") or [e.get("detail", e.get("message", "")) for e in batch.errors if e.get("row") == row["row"]]
                sheet.append([row["row"], *supplied, row["classification"], " ".join(reasons)])
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        sheet.freeze_panes = "A2"
        stream = BytesIO()
        wb.save(stream)
        response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{"sub-agent" if register_type == "SUB_AGENT" else "terminal"}-import-exclusions.xlsx"'
        return response

    @action(detail=True, methods=["post"])
    @transaction.atomic
    def cancel(self, request, pk=None):
        batch = self.locked_batch()
        batch.status = "CANCELLED"
        batch.save()
        audit(request.user, batch, "TERMINAL_IMPORT_CANCELLED")
        return Response(self.get_serializer(batch).data)
