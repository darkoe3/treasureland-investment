from datetime import date, timedelta
from io import BytesIO, StringIO
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from core.importers import parse_daily_sheet_workbook
from core.models import AuditLog, Person, TerminalImportBatch, TerminalNumber, TPMCode, TPMDailyTransaction, User
from core.tests.test_daily_sheet_imports import workbook_upload
from core.tests.test_phase_4 import Phase4Mixin


def upload(rows):
    wb = Workbook()
    wb.active.append(["S/NOS", "SUB AGT NOS", "TERMINAL NOS", "NAME"])
    for row in rows:
        wb.active.append(row)
    stream = BytesIO()
    wb.save(stream)
    return SimpleUploadedFile("register.xlsx", stream.getvalue())


class TerminalRegisterTests(Phase4Mixin, APITestCase):
    def setUp(self):
        self.setUpBase()
        self.client.force_authenticate(self.admin)
        self.base = "/api/terminal-numbers/"
        self.imports = "/api/terminal-number-imports/"

    def data(self, code=None, **kwargs):
        code = code or self.tpm_a
        return {"agency": code.person.agency_id, "person": code.person_id,
                "sub_agent_number": code.pk, "terminal_number": "00123", **kwargs}

    def create(self, **kwargs):
        response = self.client.post(self.base, self.data(**kwargs), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return TerminalNumber.objects.get(pk=response.data["id"])

    def preview(self, rows=None):
        response = self.client.post(self.imports + "preview/", {"agency": self.musa.pk,
            "file": upload(rows if rows is not None else [[1, self.tpm_a.code, "00009", "  ayo  "]])}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def confirm(self, batch):
        return self.client.post(f"{self.imports}{batch['id']}/confirm/", {"confirmed": True}, format="json")

    def reassign(self, terminal, code=None, **kwargs):
        return self.client.post(f"{self.base}{terminal.pk}/reassign/",
            self.data(code=code or self.tpm_sub, reason="Replacement allocation", confirmed=True, **kwargs), format="json")

    def test_manual_text_leading_zeroes_relationships_and_audit(self):
        terminal = self.create(terminal_number=" 001Ab ")
        self.assertEqual(terminal.terminal_number, "001Ab")
        self.assertEqual(terminal.person_id, self.tpm_a.person_id)
        self.assertEqual(AuditLog.objects.get(action="TERMINAL_CREATED").new_values["terminal_number"], "001Ab")
        self.assertEqual(TPMCode.objects.get(pk=self.tpm_a.pk).code, "TPM-A")
        for bad in (123, True, " ", "x" * 81):
            self.assertEqual(self.client.post(self.base, self.data(code=self.tpm_b, terminal_number=bad), format="json").status_code, 400)

    def test_relationship_mismatch(self):
        for attrs in ({"agency": self.sango.pk}, {"person": self.sub.pk}, {"sub_agent_number": 999999}):
            response = self.client.post(self.base, self.data(**attrs), format="json")
            self.assertEqual(response.status_code, 400)
        self.assertFalse(TerminalNumber.objects.exists())

    def test_uniqueness_and_inactive_identifier_reservation(self):
        terminal = self.create(terminal_number="00Ab")
        self.assertEqual(self.client.post(self.base, self.data(code=self.tpm_b, terminal_number="00aB"), format="json").status_code, 400)
        self.assertEqual(self.client.post(self.base, self.data(terminal_number="other"), format="json").status_code, 400)
        self.client.post(f"{self.base}{terminal.pk}/deactivate/")
        self.assertEqual(self.client.post(self.base, self.data(code=self.tpm_b, terminal_number="00AB"), format="json").status_code, 400)
        self.create(terminal_number="different")
        self.assertEqual(self.client.post(f"{self.base}{terminal.pk}/reactivate/").status_code, 400)

    def test_deactivate_reactivate_and_edit_history(self):
        terminal = self.create()
        url = f"{self.base}{terminal.pk}/"
        self.assertEqual(self.client.patch(url, {"terminal_number": "000456"}, format="json").status_code, 200)
        self.assertEqual(self.client.patch(url, {"person": self.sub.pk}, format="json").status_code, 400)
        self.assertEqual(self.client.post(url + "deactivate/").status_code, 200)
        terminal.refresh_from_db()
        self.assertEqual(terminal.deactivated_by, self.admin)
        self.assertIsNotNone(terminal.deactivated_at)
        self.assertEqual(self.client.post(url + "reactivate/").status_code, 200)
        terminal.refresh_from_db()
        self.assertIsNone(terminal.deactivated_at)
        self.assertEqual(self.client.delete(url).status_code, 405)
        self.assertEqual(len(self.client.get(url + "history/").data), 4)

    def test_accountant_isolation_and_write_restrictions(self):
        own = self.create()
        hidden = self.create(code=self.other_tpm, terminal_number="hidden")
        self.client.force_authenticate(self.acct)
        self.assertEqual([r["id"] for r in self.client.get(self.base).data], [own.id])
        self.assertEqual(self.client.get(f"{self.base}{hidden.pk}/").status_code, 404)
        self.assertEqual(self.client.get(self.base + f"?agency={self.sango.pk}").data, [])
        for action in ("deactivate", "reactivate", "reassign"):
            self.assertEqual(self.client.post(f"{self.base}{own.pk}/{action}/", {}, format="json").status_code, 403)
        self.assertEqual(self.client.post(self.base, self.data(), format="json").status_code, 403)
        self.assertEqual(self.client.patch(f"{self.base}{own.pk}/", {"terminal_number": "new"}).status_code, 403)
        self.assertEqual(self.client.post(self.imports + "preview/", {}).status_code, 403)
        self.assertEqual(self.client.get(f"{self.base}{own.pk}/history/").status_code, 403)

    def test_reassignment_requires_confirmation_reason_and_target(self):
        terminal = self.create()
        for changes in ({"confirmed": False}, {"reason": " "}, {"agency": self.sango.pk}):
            data = {**self.data(code=self.tpm_sub), "reason": "reason", "confirmed": True, **changes}
            self.assertEqual(self.client.post(f"{self.base}{terminal.pk}/reassign/", data, format="json").status_code, 400)
        self.create(code=self.tpm_sub, terminal_number="occupied")
        self.assertEqual(self.reassign(terminal).status_code, 400)

    def test_cross_agency_reassignment_snapshots_and_immutable_history(self):
        terminal = self.create()
        txn = self.txn(self.sheet(), self.tpm_a)
        before = (txn.person_name_snapshot, txn.tpm_code_snapshot, txn.terminal_number_snapshot)
        self.assertEqual(before[-1], "00123")
        self.assertEqual(self.reassign(terminal, self.other_tpm).status_code, 200)
        txn.refresh_from_db()
        self.assertEqual(before, (txn.person_name_snapshot, txn.tpm_code_snapshot, txn.terminal_number_snapshot))
        log = AuditLog.objects.get(action="TERMINAL_REASSIGNED")
        self.assertEqual(log.old_values["agency"], self.musa.pk)
        self.assertEqual(log.new_values["agency"], self.sango.pk)
        for action in (lambda: log.save(), lambda: log.delete(), lambda: AuditLog.objects.filter(pk=log.pk).update(description="rewrite"), lambda: AuditLog.objects.filter(pk=log.pk).delete()):
            with self.assertRaises(ValidationError):
                action()
        self.client.force_authenticate(self.acct)
        self.assertEqual(self.client.get(f"{self.base}{terminal.pk}/").status_code, 404)

    def test_indirect_owner_changes_are_blocked(self):
        self.create()
        self.assertEqual(self.client.patch(f"/api/tpm-codes/{self.tpm_a.pk}/", {"person": self.sub.pk, "confirm_reassignment": True}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"/api/people/{self.person.pk}/", {"agency": self.sango.pk}, format="json").status_code, 400)

    def test_template_text_no_formulas_and_round_trip(self):
        response = self.client.get(self.imports + f"template/?agency={self.musa.pk}")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual([c.value for c in wb.active[1]], ["S/NOS", "SUB AGT NOS", "TERMINAL NOS", "NAME"])
        self.assertEqual(wb.active["B2"].number_format, "@")
        self.assertEqual(wb.active["C2"].number_format, "@")
        wb.active["B2"], wb.active["C2"], wb.active["D2"] = self.tpm_a.code, "000099", self.person.full_name
        for sheet in wb:
            self.assertFalse(any(c.data_type == "f" for row in sheet for c in row))
        stream = BytesIO(); wb.save(stream)
        preview = self.client.post(self.imports + "preview/", {"agency": self.musa.pk, "file": SimpleUploadedFile("template.xlsx", stream.getvalue())}, format="multipart")
        self.assertEqual(preview.data["warnings"], [])
        self.assertEqual(preview.data["errors"], [])
        self.assertEqual(self.confirm(preview.data).status_code, 200)
        self.assertEqual(TerminalNumber.objects.get().terminal_number, "000099")

    def test_preview_has_no_terminal_writes_then_atomic_confirm(self):
        batch = self.preview()
        self.assertFalse(TerminalNumber.objects.exists())
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "ADD_TERMINAL_TO_EXISTING_SUBAGENT")
        self.assertEqual(self.confirm(batch).status_code, 200)
        self.assertEqual(TerminalNumber.objects.get().terminal_number, "00009")
        self.assertEqual(self.confirm(batch).status_code, 400)
        batch = self.preview()
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "UNCHANGED")
        self.assertEqual(self.confirm(batch).data["result_counts"], {"created": 0, "unchanged": 1})
        log = AuditLog.objects.filter(action="TERMINAL_IMPORT_PREVIEWED").first()
        self.assertNotIn("rows", log.new_values)

    def test_preview_duplicates_formulas_names_partial_rows_and_numeric_warning(self):
        rows = [[1, self.tpm_a.code, "0001", "Ayo"], [2, self.tpm_a.code.lower(), "0001", "wrong"],
                [3, self.tpm_b.code, None, "Ayo"], [4, "=1+1", "1", "Ayo"],
                [5, 123, 123, "Ayo"], [None, None, None, None]]
        batch = self.preview(rows)
        messages = " ".join(e["message"] for e in batch["errors"])
        for text in ("Row 3", "Duplicate", "required", "Formula", "does not exist"):
            self.assertIn(text, messages)
        self.assertTrue(batch["warnings"])
        self.assertEqual(len(batch["preview_payload"]["rows"]), 5)
        self.assertEqual(self.confirm(batch).status_code, 400)

    def test_conflict_classification_and_no_overwrite(self):
        terminal = self.create()
        batch = self.preview([[1, self.tpm_b.code, terminal.terminal_number, "Ayo"]])
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "REASSIGNMENT_REQUIRED")
        batch = self.preview([[1, self.tpm_a.code, "other", "Ayo"]])
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "REASSIGNMENT_REQUIRED")
        self.client.post(f"{self.base}{terminal.pk}/deactivate/")
        batch = self.preview([[1, self.tpm_a.code, terminal.terminal_number, "Ayo"]])
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "REASSIGNMENT_REQUIRED")

    def test_confirm_rolls_back_first_row_when_second_row_changed(self):
        batch = self.preview([[1, self.tpm_a.code, "new1", "Ayo"], [2, self.tpm_b.code, "new2", "Ayo"]])
        self.create(code=self.tpm_sub, terminal_number="new2")
        self.assertEqual(self.confirm(batch).status_code, 400)
        self.assertFalse(TerminalNumber.objects.filter(terminal_number="new1").exists())
        self.assertEqual(TerminalImportBatch.objects.get(pk=batch["id"]).status, "PREVIEWED")

    def test_batch_ownership_expiry_cancel_and_confirmation(self):
        batch = self.preview()
        other_admin = User.objects.create_superuser("other-admin@example.com", "AdminPass456!", full_name="Other admin")
        self.client.force_authenticate(other_admin)
        self.assertEqual(self.confirm(batch).status_code, 404)
        self.assertEqual(self.client.get(f"{self.imports}{batch['id']}/").status_code, 404)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.post(f"{self.imports}{batch['id']}/confirm/", {}).status_code, 400)
        self.assertEqual(self.client.post(f"{self.imports}{batch['id']}/cancel/").status_code, 200)
        self.assertEqual(self.confirm(batch).status_code, 400)
        self.assertTrue(AuditLog.objects.filter(action="TERMINAL_IMPORT_CANCELLED").exists())
        batch = self.preview()
        TerminalImportBatch.objects.filter(pk=batch["id"]).update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.confirm(batch).status_code, 400)

    def test_failed_audit_rolls_back_terminal(self):
        with patch("core.terminal_register.audit", side_effect=RuntimeError("unavailable")):
            with self.assertRaises(RuntimeError):
                self.client.post(self.base, self.data(), format="json")
        self.assertFalse(TerminalNumber.objects.exists())

    def test_daily_import_system_mapping_and_conflict(self):
        self.create()
        def parse(register):
            return parse_daily_sheet_workbook(workbook_upload(rows=[{"sub": self.tpm_a.code, "amounts": [10, 20]}], headers=["Premier", "Lucky"], register_rows=register, raw_date=date(2026, 8, 24)), self.musa, date(2026, 8, 24))
        result = parse([(self.tpm_a.code, "00123", "Ayo")])
        self.assertEqual(result.errors, [])
        self.assertEqual(result.payload["rows"][0]["terminal_number"], "00123")
        self.assertEqual(parse([]).errors, [])
        self.assertTrue(parse([(self.tpm_a.code, "wrong", "Ayo")]).errors)
        self.assertTrue(parse([(self.tpm_a.code, "00123", "wrong")]).errors)

    def test_report_reads_snapshot_and_blank_not_current_assignment(self):
        sheet = self.sheet()
        old = self.txn(sheet, self.tpm_a)
        terminal = self.create()
        self.txn(sheet, self.tpm_b)
        self.reassign(terminal, self.other_tpm)
        response = self.client.get(f"/api/reports/agency-summary/?agency={self.musa.pk}&period=daily&date=2026-08-24&status=DRAFT")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(all(row["terminal_number"] == "Not recorded" for row in response.data["details"]))
        old.refresh_from_db()
        self.assertEqual(old.terminal_number_snapshot, "")

    def test_readiness_command_is_read_only(self):
        self.txn(self.sheet(), self.tpm_a)
        output = StringIO()
        call_command("terminal_readiness", stdout=output)
        self.assertIn('"historical_transactions_without_terminal_snapshot": 1', output.getvalue())
        self.assertFalse(TerminalNumber.objects.exists())
        self.assertEqual(TPMDailyTransaction.objects.count(), 1)


class TerminalOnboardingTests(Phase4Mixin, APITestCase):
    setUp = TerminalRegisterTests.setUp
    data = TerminalRegisterTests.data
    create = TerminalRegisterTests.create

    def test_blank_template_rows_are_ignored_in_both_modes_with_excel_row_numbers(self):
        for mode in ("LINK_EXISTING", "ONBOARD_MISSING"):
            with self.subTest(mode=mode):
                response = self.client.post(self.imports + "preview/", {
                    "agency": self.musa.pk, "mode": mode,
                    "file": upload([[1, None, None, None], [None, None, None, None],
                                    [3, " \t", "", "\u00a0"],
                                    [4, self.tpm_a.code, "00009", "Ayo"]]),
                }, format="multipart")
                self.assertEqual(response.status_code, 201)
                self.assertEqual(response.data["errors"], [])
                self.assertEqual(response.data["warnings"], [])
                payload = response.data["preview_payload"]
                self.assertEqual(payload["ignored_blank_rows"], 3)
                self.assertEqual([row["row"] for row in payload["rows"]], [5])
                self.assertEqual(payload["creation_counts"]["terminals"], 1)

    def test_each_missing_required_field_still_blocks_after_blank_rows(self):
        for mode in ("LINK_EXISTING", "ONBOARD_MISSING"):
            with self.subTest(mode=mode):
                response = self.client.post(self.imports + "preview/", {
                    "agency": self.musa.pk, "mode": mode,
                    "file": upload([[1, None, None, None], [2, None, "t1", "Ayo"],
                                    [3, "s2", None, "Ayo"], [4, "s3", "t3", None]]),
                }, format="multipart")
                self.assertEqual(response.status_code, 201)
                payload = response.data["preview_payload"]
                self.assertEqual(payload["ignored_blank_rows"], 1)
                self.assertEqual([r["classification"] for r in payload["rows"]], ["INVALID"] * 3)
                self.assertEqual([e["row"] for e in response.data["errors"]], [3, 4, 5])
                self.assertEqual(self.approve(response.data).status_code, 400)
                self.assertFalse(TerminalNumber.objects.exists())

    def test_69_serial_only_rows_do_not_hide_six_name_conflicts(self):
        codes = [TPMCode.objects.create(person=self.person, code=f"conflict-{n}") for n in range(6)]
        before = list(TPMCode.objects.values())
        batch = self.onboard([[n, None, None, None] for n in range(1, 70)] +
                             [[70 + n, code.code, str(1000 + n), "Wrong name"] for n, code in enumerate(codes)])
        payload = batch["preview_payload"]
        self.assertEqual(payload["ignored_blank_rows"], 69)
        self.assertEqual(payload["summary"]["NAME_CONFLICT"], 6)
        self.assertEqual(payload["summary"]["INVALID"], 0)
        self.assertEqual([e["row"] for e in batch["errors"]], list(range(71, 77)))
        self.assertEqual(self.approve(batch).status_code, 400)
        self.assertEqual(list(TPMCode.objects.values()), before)
        self.assertFalse(TerminalNumber.objects.exists())

    def test_numeric_warnings_keep_excel_row_numbers_after_ignored_rows(self):
        batch = self.onboard([[1, None, None, None], [None, None, None, None], [3, 123, 456, "New"]])
        self.assertEqual(batch["preview_payload"]["ignored_blank_rows"], 2)
        self.assertEqual([w["row"] for w in batch["warnings"]], [4, 4])
        for warning in batch["warnings"]:
            self.assertTrue(warning["message"].startswith("Row 4 — "))

    def onboard(self, rows):
        response = self.client.post(self.imports + "preview/", {"agency": self.musa.pk,
            "mode": "ONBOARD_MISSING", "file": upload(rows)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def approve(self, batch, **changes):
        return self.client.post(f"{self.imports}{batch['id']}/confirm/", {
            "confirmed": True, "reason": "Verified source register", "warnings_acknowledged": True,
            "acknowledged_counts": batch["preview_payload"]["creation_counts"], **changes}, format="json")

    def test_onboarding_dependency_order_counts_and_immutable_summary(self):
        before = list(Person.objects.values())
        batch = self.onboard([[1, "001NEW", "000NEW", "  New   Person "],
                              [2, "002NEW", "000NEXT", "new person"],
                              [3, "003NEW", "000AYO", " aYo "],
                              [4, self.tpm_a.code, "000EXIST", "Ayo"]])
        self.assertEqual(batch["errors"], [])
        self.assertEqual(batch["preview_payload"]["creation_counts"], {"people": 1, "sub_agent_numbers": 3, "terminals": 4})
        self.assertEqual(list(Person.objects.values()), before)
        self.assertEqual(self.approve(batch).status_code, 200)
        self.assertEqual(TPMCode.objects.get(code="001NEW").person_id, TPMCode.objects.get(code="002NEW").person_id)
        self.assertEqual(TPMCode.objects.get(code="003NEW").person_id, self.person.pk)
        for original in before:
            self.assertEqual(Person.objects.values().get(pk=original["id"]), original)
        log = AuditLog.objects.get(action="TERMINAL_IMPORT_CONFIRMED")
        self.assertEqual(log.new_values["creation_counts"]["people"], 1)
        self.assertEqual(log.new_values["mode"], "ONBOARD_MISSING")
        self.assertEqual(log.description, "Verified source register")
        self.assertNotIn("rows", log.new_values)
        with self.assertRaises(ValidationError):
            AuditLog.objects.filter(pk=log.pk).update(description="changed")

    def test_onboarding_blocks_names_agencies_partial_duplicates_and_ignores_blanks(self):
        Person.objects.create(agency=self.musa, full_name=" AYO ", agent_type="SUBAGENT")
        batch = self.onboard([[1, "new", "t1", "ayo"], [2, self.tpm_a.code, "t2", "Wrong"],
            [3, self.other_tpm.code, "t3", self.other_tpm.person.full_name],
            [4, "partial", None, "Name"], [5, None, None, None],
            [6, "ok", "t6", "New"], [7, "OK", "t7", "New"]])
        self.assertEqual([r["classification"] for r in batch["preview_payload"]["rows"]],
            ["NAME_CONFLICT", "NAME_CONFLICT", "AGENCY_CONFLICT", "INVALID", "CREATE_PERSON_SUBAGENT_TERMINAL", "DUPLICATE"])
        self.assertEqual(self.approve(batch).status_code, 400)
        self.assertFalse(TerminalNumber.objects.exists())

    def test_onboarding_confirmation_requirements_and_accountant_denial(self):
        batch = self.onboard([[1, "new", "term", "New"]])
        for changes in ({"confirmed": False}, {"reason": " "}, {"acknowledged_counts": {}}, {"reason": None}):
            self.assertEqual(self.approve(batch, **changes).status_code, 400)
        self.client.force_authenticate(self.acct)
        self.assertEqual(self.approve(batch).status_code, 403)
        self.assertEqual(self.client.post(self.imports + "preview/", {"agency": self.musa.pk,
            "mode": "ONBOARD_MISSING", "file": upload([[1, "new", "term", "New"]])}, format="multipart").status_code, 403)

    def test_numeric_fields_and_acknowledgement(self):
        batch = self.onboard([[1, 123, 456, "New"]])
        self.assertEqual([w["field"] for w in batch["warnings"]], ["SUB AGT NOS", "TERMINAL NOS"])
        for warning in batch["warnings"]:
            self.assertEqual(warning["message"],
                f"Row 2 — {warning['field']} is numeric and may have lost leading zeroes.")
        self.assertEqual(self.approve(batch, warnings_acknowledged=False).status_code, 400)
        self.assertEqual(self.approve(batch).status_code, 200)
        self.assertEqual(TerminalNumber.objects.get().terminal_number, "456")

    def test_atomic_onboarding_failure_rolls_back_every_record_and_audit(self):
        batch = self.onboard([[1, "new1", "t1", "New One"], [2, "new2", "t2", "New Two"]])
        before = (Person.objects.count(), TPMCode.objects.count(), AuditLog.objects.count())
        from core.terminal_register import create_terminal
        def fail_second(data, user):
            if data["terminal_number"] == "t2":
                raise ValidationError("Injected failure")
            return create_terminal(data, user)
        with patch("core.terminal_register.create_terminal", side_effect=fail_second):
            self.assertEqual(self.approve(batch).status_code, 400)
        self.assertEqual((Person.objects.count(), TPMCode.objects.count(), AuditLog.objects.count()), before)
        self.assertFalse(TerminalNumber.objects.exists())
        self.assertEqual(TerminalImportBatch.objects.get(pk=batch["id"]).status, "PREVIEWED")

    def test_stale_matching_person_blocks_and_historical_records_unchanged(self):
        historic = self.txn(self.sheet(), self.tpm_a)
        before = historic.__class__.objects.values().get(pk=historic.pk)
        batch = self.onboard([[1, "new", "t1", "New"]])
        Person.objects.create(agency=self.musa, full_name="New", agent_type="SUBAGENT")
        self.assertEqual(self.approve(batch).status_code, 400)
        self.assertFalse(TPMCode.objects.filter(code="new").exists())
        batch = self.onboard([[1, "new", "t1", "New"]])
        self.assertEqual(self.approve(batch).status_code, 200)
        self.assertEqual(historic.__class__.objects.values().get(pk=historic.pk), before)

    def test_onboarding_never_reassigns_existing_terminal(self):
        terminal = self.create()
        before = TerminalNumber.objects.values().get(pk=terminal.pk)
        batch = self.onboard([[1, "new", terminal.terminal_number, "New"]])
        self.assertEqual(batch["preview_payload"]["rows"][0]["classification"], "REASSIGNMENT_REQUIRED")
        self.assertEqual(self.approve(batch).status_code, 400)
        self.assertEqual(TerminalNumber.objects.values().get(pk=terminal.pk), before)

    def test_link_default_does_not_onboard_and_unknown_modes_are_rejected(self):
        for mode in (None, 'LINK_EXISTING', 'automatic'):
            data = {'agency': self.musa.pk, 'file': upload([[1, 'new', 't1', 'New']])}
            if mode is not None:
                data['mode'] = mode
            response = self.client.post(self.imports + 'preview/', data, format='multipart')
            if mode == 'automatic':
                self.assertEqual(response.status_code, 400)
            else:
                self.assertEqual(response.status_code, 201)
                self.assertTrue(response.data['errors'])
                self.assertEqual(response.data['preview_payload']['mode'], 'LINK_EXISTING')
                self.assertEqual(self.approve(response.data).status_code, 400)
        self.assertFalse(Person.objects.filter(full_name='New').exists())
        self.assertFalse(TPMCode.objects.filter(code='new').exists())

    def test_foreign_terminal_identifier_blocks_onboarding(self):
        terminal = self.create(code=self.other_tpm, terminal_number='foreign')
        batch = self.onboard([[1, 'new', terminal.terminal_number, 'New']])
        self.assertEqual(batch['preview_payload']['rows'][0]['classification'], 'AGENCY_CONFLICT')
        self.assertEqual(self.approve(batch).status_code, 400)

    def test_inactive_unique_person_and_empty_workbook_block(self):
        self.person.is_active = False
        self.person.save()
        for rows in ([[1, 'new', 't1', 'Ayo']], [[1, None, None, None]]):
            batch = self.onboard(rows)
            self.assertTrue(batch['errors'])
            self.assertEqual(self.approve(batch).status_code, 400)

    def test_final_audit_failure_rolls_back_onboarding(self):
        batch = self.onboard([[1, 'new', 't1', 'New']])
        from core.terminal_register import audit
        def fail_final(user, obj, action_name, *args, **kwargs):
            if action_name == 'TERMINAL_IMPORT_CONFIRMED':
                raise RuntimeError('Audit unavailable')
            return audit(user, obj, action_name, *args, **kwargs)
        with patch('core.terminal_register.audit', side_effect=fail_final):
            with self.assertRaises(RuntimeError):
                self.approve(batch)
        self.assertFalse(Person.objects.filter(full_name='New').exists())
        self.assertFalse(TPMCode.objects.filter(code='new').exists())
        self.assertFalse(TerminalNumber.objects.exists())
        self.assertFalse(AuditLog.objects.filter(action='TERMINAL_CREATED').exists())
